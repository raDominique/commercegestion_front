#!/usr/bin/env python3
"""Convertit un fichier Markdown de docs/ en document Word (.docx).

Usage :
    python convert_to_docx.py [fichier.md] [fichier.docx]

D\u00e9fauts (chemins relatifs au dossier du script, c.-\u00e0-d. docs/) :
    entr\u00e9e : DOCUMENTATION_FRONTEND.md
    sortie : Documentation.docx

Exemples :
    python convert_to_docx.py
    python convert_to_docx.py "Guide utilisateur.md" "Guide utilisateur.docx"
    python convert_to_docx.py ../README.md /tmp/readme.docx

Supporte : titres, paragraphes, listes, tableaux, blocs de code, images
Markdown (chemins relatifs au fichier .md), `code inline` et **gras**.
"""

import os
import re
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE = os.path.dirname(os.path.abspath(__file__))

CODE_FONT = "Consolas"
INLINE_CODE_COLOR = RGBColor(0x6D, 0x28, 0xD9)   # violet
MUTED_COLOR = RGBColor(0x6B, 0x72, 0x80)         # gris


def resolve_path(arg, default):
    """Chemin fourni par l'utilisateur (cwd ou relatif au script) ou défaut."""
    if not arg:
        return os.path.join(BASE, default)
    if os.path.isabs(arg):
        return arg
    if os.path.exists(arg):
        return os.path.abspath(arg)
    return os.path.join(BASE, arg)


# --------------------------------------------------------------------- #
# Formatage en ligne : `code` et **gras**
# --------------------------------------------------------------------- #

def add_rich_text(paragraph, text, base_size=11, base_bold=False):
    """Ajoute du texte en gérant `code` et **gras**."""
    # découpage alterné autour des backticks
    parts = re.split(r'(`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if len(part) > 2 and part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = INLINE_CODE_COLOR
            continue
        # gras
        for sub in re.split(r'(\*\*[^*]+\*\*)', part):
            if not sub:
                continue
            if sub.startswith('**') and sub.endswith('**') and len(sub) > 4:
                run = paragraph.add_run(sub[2:-2])
                run.bold = True
            else:
                run = paragraph.add_run(sub)
            run.font.size = Pt(base_size)
            if base_bold:
                run.bold = True


def add_paragraph(doc, text, style=None, size=11, italic=False, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    add_rich_text(p, text, base_size=size)
    if italic or color is not None:
        for run in p.runs:
            if italic:
                run.italic = True
            if color is not None:
                run.font.color.rgb = color
    return p


# --------------------------------------------------------------------- #
# Éléments structurés
# --------------------------------------------------------------------- #

def add_heading(doc, title, md_level):
    """# → Titre, ## → niveau 1, ### → niveau 2, etc. (max 4)."""
    if md_level == 1:
        heading = doc.add_heading(title, level=0)
    else:
        heading = doc.add_heading(title, level=min(md_level - 1, 4))
    heading.paragraph_format.space_after = Pt(12)
    heading.paragraph_format.space_before = Pt(6)
    return heading


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("\u2500" * 70)
    run.font.size = Pt(4)
    run.font.color.rgb = RGBColor(128, 128, 128)


def add_code_block(doc, code):
    for idx, line in enumerate(code.split('\n')):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0 if idx < len(code.split('\n')) - 1 else 6)
        p.paragraph_format.space_before = Pt(0 if idx == 0 else 0)
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(line if line else ' ')
        run.font.name = CODE_FONT
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)


def put_picture(paragraph, alt, src, base_dir, width_in=6.0):
    """Insère une image dans un paragraphe existant. Retourne True si OK."""
    path = src if os.path.isabs(src) else os.path.normpath(os.path.join(base_dir, src))
    if os.path.exists(path):
        try:
            paragraph.add_run().add_picture(path, width=Inches(width_in))
            return True
        except Exception as exc:  # image illisible
            run = paragraph.add_run(f"[Image illisible : {src} \u2014 {exc}]")
            run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
            return False
    run = paragraph.add_run(f"[Image introuvable : {src}]")
    run.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
    return False


def add_image(doc, alt, src, base_dir):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    put_picture(p, alt, src, base_dir, width_in=6.0)
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        run = cap.add_run(alt)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED_COLOR


def add_table(doc, rows, base_dir):
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row_data in enumerate(rows):
        for j in range(num_cols):
            cell_text = row_data[j] if j < len(row_data) else ''
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            # images éventuelles dans la cellule : ![alt](src)
            parts = re.split(r'(!\[[^\]]*\]\([^)]+\))', cell_text)
            for part in parts:
                if not part:
                    continue
                m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)$', part)
                if m:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    put_picture(p, m.group(1), m.group(2), base_dir, width_in=2.2)
                else:
                    add_rich_text(p, part, base_size=10, base_bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def is_table_separator(line):
    return bool(re.match(r'^\|[\s:\-|]+\|$', line.strip()))


def is_table_line(line):
    s = line.strip()
    return s.startswith('|') and s.endswith('|') and len(s) > 1


# --------------------------------------------------------------------- #
# Conversion
# --------------------------------------------------------------------- #

def convert(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    base_dir = os.path.dirname(os.path.abspath(md_path))
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # bloc de code ``` ... ```
        if stripped.startswith('```'):
            j = i + 1
            buf = []
            while j < len(lines) and not lines[j].strip().startswith('```'):
                buf.append(lines[j])
                j += 1
            add_code_block(doc, '\n'.join(buf))
            i = j + 1
            continue

        # image seule : ![alt](src)
        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', stripped)
        if m:
            add_image(doc, m.group(1), m.group(2), base_dir)
            i += 1
            continue

        # titre
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            add_heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue

        # filet horizontal
        if re.match(r'^\s*(-{3,}|\*{3,}|_{3,})\s*$', line):
            add_rule(doc)
            i += 1
            continue

        # tableau
        if is_table_line(line):
            rows = []
            while i < len(lines) and is_table_line(lines[i]):
                raw = lines[i].strip()
                if not is_table_separator(raw):
                    cells = [c.strip() for c in raw.strip('|').split('|')]
                    rows.append(cells)
                i += 1
            add_table(doc, rows, base_dir)
            continue

        # liste à puces / numérotée
        m = re.match(r'^(\s*)[-*+]\s+(.*)$', line)
        if m:
            add_paragraph(doc, m.group(2), style='List Bullet')
            i += 1
            continue
        m = re.match(r'^(\s*)\d+\.\s+(.*)$', line)
        if m:
            add_paragraph(doc, m.group(2), style='List Number')
            i += 1
            continue

        # citation
        if stripped.startswith('>'):
            add_paragraph(doc, stripped.lstrip('>').strip(), italic=True, color=MUTED_COLOR)
            i += 1
            continue

        # ligne vide : simple séparation (gérée par space_after)
        if not stripped:
            i += 1
            continue

        # paragraphe simple
        add_paragraph(doc, line)
        i += 1

    doc.save(docx_path)
    return docx_path


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('-')]
    if not args:
        md_path = resolve_path(None, 'DOCUMENTATION_FRONTEND.md')
        docx_path = resolve_path(None, 'Documentation.docx')
    else:
        md_path = resolve_path(args[0], 'DOCUMENTATION_FRONTEND.md')
        if len(args) > 1:
            docx_path = resolve_path(args[1], 'Documentation.docx')
        else:
            docx_path = os.path.splitext(md_path)[0] + '.docx'

    if not os.path.exists(md_path):
        print(f"Fichier Markdown introuvable : {md_path}")
        sys.exit(1)

    convert(md_path, docx_path)
    size = os.path.getsize(docx_path)
    print(f"Document Word créé avec succès : {docx_path} ({size} octets)")


if __name__ == '__main__':
    main()
